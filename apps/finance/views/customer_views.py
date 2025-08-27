from django.shortcuts import render, get_object_or_404, redirect
from django.http import JsonResponse, HttpResponse
from django.views import View
from django.db import transaction
from django.db.models import Q, Sum
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from django.contrib import messages
from ..models import (
    Customer, CustomerAccount, CustomerAccountBalances, Currency,
    Invoice, Payment, CustomerDeposits, Account, AccountBalance,
    Cashbook
)
from ..forms import CustomerForm, customerDepositsForm, customerDepositsRefundForm
from ..tasks import send_account_statement_email
from loguru import logger
import json
import openpyxl
from openpyxl.styles import Font, Alignment
from decimal import Decimal
import re
import datetime

@login_required
def customer(request):
    if request.method == 'GET':
        customers = Customer.objects.all().values()
        return JsonResponse(list(customers), safe=False)

    elif request.method == 'POST':
        data = json.loads(request.body)

        if Customer.objects.filter(phone_number=data['phonenumber']).exists():
            return JsonResponse({'success': False, 'message': 'Customer exists'})
        else:
            customer = Customer.objects.create(
                name=data['name'],
                email=data['email'],
                address=data['address'],
                phone_number=data['phonenumber'],
                branch=request.user.branch
            )
            account = CustomerAccount.objects.create(customer=customer)

            balances_to_create = [
                CustomerAccountBalances(account=account, currency=currency, balance=0)
                for currency in Currency.objects.all()
            ]
            CustomerAccountBalances.objects.bulk_create(balances_to_create)

        return JsonResponse({'success': True, 'message': 'Customer successfully created'})

    else:
        return JsonResponse({'success': False, 'message': 'Invalid request method'})

def validate_customer_data(data):
    errors = {}
    if 'name' not in data or len(data['name']) < 2:
        errors['name'] = 'Name is required and must be at least 2 characters long.'

    if 'email' in data and data['email'] and not validate_email_address(data['email']):
        errors['email'] = 'A valid email address is required.'

    if 'address' not in data:
        errors['address'] = 'Address is required.'

    if 'phonenumber' not in data:
        errors['phonenumber'] = 'Phone number is required.'

    return errors

def validate_email_address(email):
    email_regex = r"^[a-zA-Z0-9.!#$%&'*+/=?^_`{|}~-]+@[a-zA-Z0-9-]+(?:\.[a-zA-Z0-9-]+)*$"
    return bool(re.match(email_regex, email))

@login_required
def customer_list(request):
    search_query = request.GET.get('q', '')

    customers = Customer.objects.filter(branch=request.user.branch)
    accounts = CustomerAccountBalances.objects.all()

    total_balances_per_currency = CustomerAccountBalances.objects.filter(account__customer__branch=request.user.branch).values('currency__name').annotate(
        total_balance=Sum('balance')
    )

    if search_query:
        customers = customers.filter(Q(name__icontains=search_query) | Q(phone_number__icontains=search_query) | Q(email__icontains=search_query))

    if 'receivable' in request.GET:
        customers = customers.filter(customeraccount__balances__balance__lt=0).distinct()

    if 'download' in request.GET:
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=customers.xlsx'

        workbook = openpyxl.Workbook()
        worksheet = workbook.active

        header_font = Font(bold=True)
        header_alignment = Alignment(horizontal='center')
        for col_num, header_title in enumerate(['Customer Name', 'Phone Number', 'Email', 'Account Balance'], start=1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.value = header_title
            cell.font = header_font
            cell.alignment = header_alignment

            column_letter = openpyxl.utils.get_column_letter(col_num)
            worksheet.column_dimensions[column_letter].width = max(len(header_title), 20)

        for customer_obj in customers:
            balance = CustomerAccountBalances.objects.filter(account__customer=customer_obj).aggregate(Sum('balance'))['balance__sum'] or 0
            worksheet.append([customer_obj.name, customer_obj.phone_number, customer_obj.email, balance])

        workbook.save(response)
        return response

    return render(request, 'customers/customers.html', {
        'customers':customers,
        'accounts':accounts,
        'total_balances_per_currency':total_balances_per_currency,
    })

@login_required
def update_customer(request, customer_id):
    customer = get_object_or_404(Customer, pk=customer_id)

    if request.method == 'POST':
        form = CustomerForm(request.POST, instance=customer)
        if form.is_valid():
            form.save()
            messages.success(request, f'{customer.name} details updated successfully')
            return redirect('finance:customer_list')
    else:
        form = CustomerForm(instance=customer)

    return render(request, 'customers/update_customer.html', {'form': form, 'customer': customer})

@login_required
def delete_customer(request, customer_id):
    if request.method == 'DELETE':
        customer = get_object_or_404(Customer, pk=customer_id)
        customer.delete()
        messages.success(request, f'Customer {customer.name} deleted successfully.')
        return JsonResponse({'status': 'success'})
    else:
        return JsonResponse({'status': 'error', 'message': 'Invalid request method.'})

@login_required
def customer_account(request, customer_id):
    form = customerDepositsForm()
    refund_form = customerDepositsRefundForm()
    customer = get_object_or_404(Customer, id=customer_id)

    account = CustomerAccountBalances.objects.filter(account__customer=customer)
    invoices = Invoice.objects.filter(customer=customer, branch=request.user.branch, status=True)
    invoice_payments = Payment.objects.filter(invoice__branch=request.user.branch, invoice__customer=customer).order_by('-payment_date')

    filters = Q()
    if request.GET.get('q'):
        filters &= Q(payment_status=request.GET['q'])
    if request.GET.get('search_query'):
        search_query = request.GET['search_query']
        filters &= (Q(invoice_number__icontains=search_query) | Q(issue_date__icontains=search_query))

    invoices = invoices.filter(filters)

    if request.GET.get('email_bool'):
        send_account_statement_email.delay(customer.id, request.user.branch.id, request.user.id)
        return JsonResponse({'message': 'Email sent'})

    return render(request, 'customer.html', {
        'form':form,
        'account': account,
        'invoices': invoices,
        'customer': customer,
        'refund_form':refund_form,
        'invoice_count': invoices.count(),
        'invoice_payments': invoice_payments,
        'paid': invoices.filter(payment_status='Paid').count(),
        'due': invoices.filter(payment_status='Partial').count(),
    })

@login_required
@transaction.atomic
def add_customer_deposit(request, customer_id):
    try:
        data = json.loads(request.body)
        amount = Decimal(data.get('amount'))
        currency_id = data.get('currency')
        payment_method = data.get('payment_method')
        reason = data.get('reason')
        payment_reference = data.get('payment_reference')

        if CustomerDeposits.objects.filter(payment_reference=payment_reference).exists():
            return JsonResponse({'success':False, 'message': f'Payment reference: {payment_reference} exists'})

        currency = Currency.objects.get(id=currency_id)

        account_types = {
            'cash': Account.AccountType.CASH,
            'bank': Account.AccountType.BANK,
            'ecocash': Account.AccountType.ECOCASH,
        }

        account_name = f"{request.user.branch} {currency.name} {payment_method.capitalize()} Account"
        account, _ = Account.objects.get_or_create(name=account_name, type=account_types[payment_method])

        account_balance, _ = AccountBalance.objects.get_or_create(
            account=account, currency=currency, branch=request.user.branch, defaults={'balance': 0}
        )

        account_balance.balance += amount
        account_balance.save()

        customer = get_object_or_404(Customer, id=customer_id)
        customer_account = CustomerAccount.objects.get(customer=customer)

        customer_account_bal_object, _ = CustomerAccountBalances.objects.get_or_create(
                account=customer_account, currency=currency, defaults={'balance': 0}
        )

        customer_deposit = CustomerDeposits.objects.create(
            customer_account=customer_account_bal_object, amount=amount, currency=currency,
            payment_method=payment_method, reason=reason, payment_reference=payment_reference,
            cashier=request.user, branch=request.user.branch
        )

        customer_account_bal_object.balance += amount
        customer_account_bal_object.save()

        Cashbook.objects.create(
            issue_date=customer_deposit.date_created,
            description=f'{customer_deposit.payment_method.upper()} deposit ({customer_deposit.customer_account.account.customer.name})',
            debit=True, credit=False, amount=customer_deposit.amount,
            currency=customer_deposit.currency, branch=customer_deposit.branch
        )

        return JsonResponse({"success":True, "message": f"Customer Deposit of {currency} {amount:2f} has been successfull"}, status=200)
    except Exception as e:
        return JsonResponse({"message": f"{e}", 'success':False}, status=500)

@login_required
def deposits_list(request):
    deposits = CustomerDeposits.objects.filter(branch=request.user.branch).order_by('-date_created')
    return render(request, 'deposits.html', {
        'deposits':deposits,
        'total_deposits': deposits.aggregate(Sum('amount'))['amount__sum'] or 0,
    })

@login_required
@transaction.atomic
def refund_customer_deposit(request, deposit_id):
    try:
        deposit = CustomerDeposits.objects.get(id=deposit_id)
    except CustomerDeposits.DoesNotExist:
        return JsonResponse({'success': False, 'message': 'Deposit not found'}, status=404)

    try:
        data = json.loads(request.body)
        amount = Decimal(data.get('amount', 0))
        if amount <= 0:
            return JsonResponse({'success': False, 'message': 'Invalid amount'}, status=400)
    except (json.JSONDecodeError, TypeError, ValueError):
        return JsonResponse({'success': False, 'message': 'Invalid input data'}, status=400)

    account_types = {
        'cash': Account.AccountType.CASH,
        'bank': Account.AccountType.BANK,
        'ecocash': Account.AccountType.ECOCASH,
    }

    account_name = f"{request.user.branch} {deposit.currency.name} {deposit.payment_method.capitalize()} Account"

    try:
        account = Account.objects.get(name=account_name, type=account_types[deposit.payment_method])
        account_balance = AccountBalance.objects.get(account=account, currency=deposit.currency, branch=request.user.branch)
    except (Account.DoesNotExist, AccountBalance.DoesNotExist) as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)

    if amount > deposit.amount:
        return JsonResponse({'success': False, 'message': 'Refund amount exceeds deposit amount'}, status=400)

    account_balance.balance -= amount
    diff_amount = deposit.amount - amount

    if diff_amount == 0:
        deposit.delete()
    else:
        deposit.amount = diff_amount
        deposit.save()

    Cashbook.objects.create(
        issue_date=datetime.date.today(),
        description=f'{deposit.payment_method.upper()} deposit refund ({deposit.customer_account.account.customer.name})',
        debit=False, credit=True, amount=amount,
        currency=deposit.currency, branch=deposit.branch
    )

    account_balance.save()
    return JsonResponse({'success': True}, status=200)

@login_required
@transaction.atomic
def edit_customer_deposit(request, deposit_id):
    try:
        deposit = CustomerDeposits.objects.get(id=deposit_id)
    except CustomerDeposits.DoesNotExist:
        messages.warning(request, 'Deposit not found')
        return redirect('finance:customer_account', deposit.customer_account.account.customer.id)

    if request.method == 'POST':
        form = customerDepositsForm(request.POST)
        if not form.is_valid():
            messages.warning(request, 'Invalid form submission')
            return redirect('finance:edit_customer_deposit', deposit_id)

        amount = Decimal(request.POST.get('amount'))
        if amount <= 0:
            messages.warning(request, 'Amount cannot be zero or negative')
            return redirect('finance:edit_customer_deposit', deposit_id)

        account_types = {
            'cash': Account.AccountType.CASH,
            'bank': Account.AccountType.BANK,
            'ecocash': Account.AccountType.ECOCASH,
        }

        account_name = f"{request.user.branch} {deposit.currency.name} {deposit.payment_method.capitalize()} Account"

        try:
            account = Account.objects.get(name=account_name, type=account_types[deposit.payment_method])
            account_balance = AccountBalance.objects.get(
                account=account, currency=deposit.currency, branch=request.user.branch
            )
        except (Account.DoesNotExist, AccountBalance.DoesNotExist) as e:
            messages.warning(request, str(e))
            return redirect('finance:edit_customer_deposit', deposit_id)

        adj_amount = amount - deposit.amount

        if adj_amount != 0:
            debit, credit = (True, False) if adj_amount > 0 else (False, True)
            account_balance.balance += adj_amount

            Cashbook.objects.create(
                issue_date=datetime.date.today(),
                description=f'{deposit.payment_method.upper()} deposit adjustment ({deposit.customer_account.account.customer.name})',
                debit=debit, credit=credit, amount=abs(adj_amount),
                currency=deposit.currency, branch=deposit.branch
            )

            account_balance.save()
            deposit.amount = amount
            deposit.save()
            messages.success(request, 'Customer deposit successfully updated')
            return redirect('finance:customer_account', deposit.customer_account.account.customer.id)
    else:
        form = customerDepositsForm(instance=deposit)

    return render(request, 'customers/edit_deposit.html', {'form': form})

@login_required
def customer_deposits(request):
    customer_id = request.GET.get('customer_id')

    if customer_id:
        deposits = CustomerDeposits.objects.filter(branch=request.user.branch, customer_account__account__customer__id=customer_id).values(
            'id', 'date_created', 'amount', 'reason', 'currency__name', 'currency__symbol',
            'payment_method', 'payment_reference', 'cashier__username'
        ).order_by('-date_created')
        return JsonResponse(list(deposits), safe=False)
    else:
        return JsonResponse({'success':False, 'message':f'{customer_id} was not provided'})

@login_required
def customer_account_transactions_json(request):
    customer_id = request.GET.get('customer_id')
    transaction_type = request.GET.get('type')

    customer = get_object_or_404(Customer, id=customer_id)

    if transaction_type == 'invoices':
        invoices = Invoice.objects.filter(
            customer=customer, branch=request.user.branch, status=True
        ).order_by('-issue_date').values(
            'issue_date', 'invoice_number', 'products_purchased', 'amount_paid',
            'amount_due', 'amount', 'user__username', 'payment_status'
        )
        return JsonResponse(list(invoices), safe=False)
    else:
        return JsonResponse({'message': 'Invalid transaction type.'}, status=400)

@login_required
def customer_account_payments_json(request):
    customer_id = request.GET.get('customer_id')
    transaction_type = request.GET.get('type')

    customer = get_object_or_404(Customer, id=customer_id)

    if transaction_type == 'invoice_payments':
        invoice_payments = Payment.objects.select_related('invoice', 'invoice__currency', 'user').filter(
            invoice__branch=request.user.branch, invoice__customer=customer
        ).order_by('-payment_date').values(
            'invoice__products_purchased', 'payment_date', 'invoice__invoice_number',
            'invoice__currency__symbol', 'invoice__payment_status', 'invoice__amount_due',
            'invoice__amount', 'user__username', 'amount_paid', 'amount_due'
        )
        return JsonResponse(list(invoice_payments), safe=False)
    else:
        return JsonResponse({'message': 'Invalid transaction type.'}, status=400)

@login_required
def customer_account_json(request, customer_id):
    account = CustomerAccountBalances.objects.filter(account__customer__id=customer_id).values(
        'currency__symbol', 'balance'
    )
    return JsonResponse(list(account), safe=False)

@login_required
def print_account_statement(request, customer_id):
    try:
        customer = get_object_or_404(Customer, id=customer_id)
        account = CustomerAccountBalances.objects.filter(account__customer=customer)
        invoices = Invoice.objects.filter(customer=customer, branch=request.user.branch, status=True)
    except:
        messages.warning(request, 'Error in processing the request')
        return redirect('finance:customer_list')

    invoice_payments = Payment.objects.select_related('invoice', 'invoice__currency', 'user').filter(
        invoice__branch=request.user.branch, invoice__customer=customer
    ).order_by('-payment_date')

    return render(request, 'customers/print_customer_statement.html', {
        'customer':customer,
        'account':account,
        'invoices':invoices,
        'invoice_payments':invoice_payments
    })
